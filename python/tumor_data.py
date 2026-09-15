"""Read tumour-volume time series from the manuscript Source Data workbook."""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook


DATA_PATH = Path(__file__).resolve().parents[1] / "data" / "tumor_growth_data.xlsx"


def load_tumor_volumes(path=DATA_PATH):
    """Return time and tumour-volume series from the workbook's first sheet.

    Column B contains time in days. Untreated and L/G-treated tumours occupy
    repeated three-column blocks (short diameter, long diameter, volume).
    """
    sheet = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = range(4, 25)
    time = np.array([sheet.cell(row, 2).value for row in rows], dtype=float)

    def volume_series(columns):
        return [
            np.array(
                [
                    np.nan if sheet.cell(row, column).value is None
                    else sheet.cell(row, column).value
                    for row in rows
                ],
                dtype=float,
            )
            for column in columns
        ]

    untreated = volume_series(range(5, 27, 3))   # 8 tumours
    treated = volume_series(range(29, 63, 3))    # 12 tumours
    return time, untreated, treated
